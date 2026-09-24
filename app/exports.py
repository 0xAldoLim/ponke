import csv
import io
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.database import Account, Category, Decision, Transaction
from app.portfolio import portfolio_snapshot
from app.validation import aware


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


async def export_finance(db, user_id, format="xlsx", timezone="Asia/Jakarta"):
    transactions = (
        await db.scalars(select(Transaction).where(Transaction.user_id == user_id).order_by(Transaction.date))
    ).all()
    header = [
        "ID",
        "Date (UTC)",
        "Amount",
        "Currency",
        "Merchant",
        "Description",
        "Category",
        "Subcategory",
        "Type",
        "Account ID",
        "Payment Method",
        "Source",
        "Source Message",
        "Original Input",
    ]
    rows = [
        [
            t.id,
            aware(t.date).isoformat(),
            t.amount,
            t.currency,
            t.merchant,
            t.description,
            t.category,
            t.subcategory,
            t.transaction_type,
            t.account_id,
            t.payment_method,
            t.source,
            t.source_message_id,
            t.original_input,
        ]
        for t in transactions
    ]
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(header)
        writer.writerows([[safe_cell(value) for value in row] for row in rows])
        return output.getvalue().encode("utf-8-sig"), "ponke-finance.csv"
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name, headings, data):
        ws = workbook.create_sheet(name)
        ws.append(headings)
        for row in data:
            ws.append([safe_cell(value) for value in row])
        return ws

    dashboard = sheet(
        "Dashboard",
        ["PONKE / FINANCE", "Value"],
        [
            ["Source", "Recorded transactions; currencies are separate"],
            ["Coverage", "Account balances are manual snapshots"],
            ["Transactions", len(transactions)],
            ["Net worth", "Reported asset and liability snapshots; no FX conversion"],
        ],
    )
    sheet("Transactions", header, rows)
    categories = (await db.scalars(select(Category).where(Category.user_id == user_id))).all()
    sheet("Categories", ["Category", "Subcategory"], [[c.name, c.subcategory] for c in categories])
    groups = {}
    for t in transactions:
        key = (aware(t.date).astimezone(ZoneInfo(timezone)).strftime("%Y-%m"), t.currency)
        values = groups.setdefault(key, {"income": Decimal(0), "expense": Decimal(0)})
        if t.transaction_type in {"income", "dividend", "interest"}:
            values["income"] += t.amount
        elif t.transaction_type == "expense":
            values["expense"] += t.amount
    monthly = sheet(
        "Monthly Summary",
        [f"Month ({timezone})", "Currency", "Income", "Expense", "Surplus", "Savings rate"],
        [],
    )
    for index, ((month, currency), values) in enumerate(sorted(groups.items()), start=2):
        monthly.append(
            [
                month,
                currency,
                values["income"],
                values["expense"],
                f"=C{index}-D{index}",
                f'=IF(C{index}=0,"",E{index}/C{index})',
            ]
        )
        monthly.cell(index, 6).number_format = "0.0%"
    accounts = (await db.scalars(select(Account).where(Account.user_id == user_id))).all()
    sheet(
        "Accounts",
        ["Name", "Type", "Currency", "Balance snapshot", "As of", "Active"],
        [
            [a.name, a.type, a.currency, a.current_balance, str(a.balance_as_of or ""), a.active]
            for a in accounts
        ],
    )
    portfolio = await portfolio_snapshot(db, user_id)
    sheet(
        "Investments",
        ["Symbol", "Quantity", "Currency", "Manual value", "Price as of", "Stale"],
        [
            [h["symbol"], h["quantity"], h["currency"], h["value"], h["price_as_of"], h["stale"]]
            for h in portfolio["holdings"]
        ],
    )
    sheet(
        "Net Worth",
        ["Currency", "Reported net worth", "Completeness"],
        [
            [
                c,
                v,
                "Partial: subtracts recorded liabilities; missing records and stale snapshots remain unknown",
            ]
            for c, v in portfolio["reported_net_worth"].items()
        ],
    )
    decisions = (await db.scalars(select(Decision).where(Decision.user_id == user_id))).all()
    sheet(
        "Decision History",
        ["ID", "Date", "Question", "Type", "Recommendation", "Confidence", "Status"],
        [
            [
                d.id,
                str(d.created_at),
                d.user_question,
                d.decision_type,
                d.final_recommendation["recommended_action"],
                d.final_confidence or "",
                d.status,
            ]
            for d in decisions
        ],
    )
    dashboard["B4"] = "=COUNTA(Transactions!A:A)-1"
    for ws in workbook:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="123B36")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EDF5F2")
                if isinstance(cell.value, (float, Decimal)):
                    cell.number_format = "#,##0.00"
        for index in range(1, ws.max_column + 1):
            width = min(
                55,
                max(
                    15,
                    max(
                        len(str(ws.cell(row, index).value or ""))
                        for row in range(1, min(ws.max_row + 1, 100))
                    )
                    + 2,
                ),
            )
            ws.column_dimensions[get_column_letter(index)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), "ponke-finance.xlsx"
