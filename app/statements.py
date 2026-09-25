"""Conservative bank-statement import with preview and idempotent commit."""

import csv
import hashlib
import io
import re
from datetime import UTC, datetime, timedelta
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path
from typing import Protocol

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app.database import Account, StatementImport, Transaction
from app.sheets import enqueue_sync
from app.validation import Clarification

DATE_NAMES = {"date", "transaction date", "tanggal", "tgl", "posting date"}
DESCRIPTION_NAMES = {"description", "details", "merchant", "keterangan", "narrative", "transaction"}
AMOUNT_NAMES = {"amount", "jumlah", "nominal", "value"}
DEBIT_NAMES = {"debit", "withdrawal", "keluar"}
CREDIT_NAMES = {"credit", "deposit", "masuk"}
CURRENCY_NAMES = {"currency", "ccy", "mata uang"}


class StatementParser(Protocol):
    def parse(self, content: bytes) -> list[dict]: ...


def parse_amount(raw):
    if raw is None:
        return None
    value = str(raw).strip().replace("Rp", "").replace("IDR", "").replace("$", "")
    if not value:
        return None
    negative = value.startswith("-") or (value.startswith("(") and value.endswith(")"))
    value = value.strip("-() ")
    if "," in value and "." in value:
        value = (
            value.replace(".", "").replace(",", ".")
            if value.rfind(",") > value.rfind(".")
            else value.replace(",", "")
        )
    elif "," in value:
        value = value.replace(",", "") if len(value.rsplit(",", 1)[1]) == 3 else value.replace(",", ".")
    try:
        result = Decimal(value)
    except InvalidOperation:
        return None
    return -result if negative else result


def parse_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def normalize_row(row, index):
    values = {str(k).strip().casefold(): v for k, v in row.items() if k is not None}

    def first(names):
        return next((values[k] for k in names if k in values and values[k] not in (None, "")), None)

    date = parse_date(first(DATE_NAMES))
    description = str(first(DESCRIPTION_NAMES) or "").strip()[:200]
    debit, credit = parse_amount(first(DEBIT_NAMES)), parse_amount(first(CREDIT_NAMES))
    raw_amount = parse_amount(first(AMOUNT_NAMES))
    if debit is not None and credit is not None and debit != 0 and credit != 0:
        return {"index": index, "status": "review", "reason": "Both debit and credit are populated."}
    if (
        debit is None
        and credit is None
        and raw_amount is not None
        and not str(first(AMOUNT_NAMES)).strip().startswith(("-", "+", "("))
    ):
        return {
            "index": index,
            "status": "review",
            "reason": "Unsigned amount has no debit/credit direction.",
        }
    amount = (
        -abs(debit)
        if debit is not None and debit != 0
        else abs(credit)
        if credit is not None and credit != 0
        else raw_amount
    )
    currency = str(first(CURRENCY_NAMES) or "IDR").strip().upper()
    if (
        not date
        or not description
        or amount is None
        or amount == 0
        or currency not in {"IDR", "USD", "EUR", "MYR", "SGD"}
    ):
        return {
            "index": index,
            "status": "review",
            "reason": "Missing or ambiguous date, description, amount, or currency.",
        }
    if not amount.is_finite() or abs(amount) >= Decimal("1e18"):
        return {"index": index, "status": "review", "reason": "Invalid amount."}
    return {
        "index": index,
        "status": "ready",
        "date": date,
        "description": description,
        "amount": str(abs(amount)),
        "currency": currency,
        "transaction_type": "expense" if amount < 0 else "income",
    }


class CSVParser:
    def parse(self, content):
        text = content.decode("utf-8-sig", errors="replace")
        if "\ufffd" in text:
            text = content.decode("cp1252", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        return list(csv.DictReader(io.StringIO(text), dialect=dialect))


class XLSXParser:
    def parse(self, content):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        header = next(values, None)
        if not header:
            return []
        return [
            dict(zip(header, row, strict=False)) for row in values if any(cell is not None for cell in row)
        ]


class PDFParser:
    def parse(self, content):
        reader = PdfReader(io.BytesIO(content), strict=False)
        if len(reader.pages) > 30:
            raise Clarification("PDF has too many pages for a safe statement preview.")
        rows = []
        pattern = re.compile(r"^(\d{4}-\d{2}-\d{2}|\d{2}/\d{2}/\d{4})\s+(.+?)\s+(-?[\d.,]+)\s*$")
        for page in reader.pages:
            for line in (page.extract_text() or "").splitlines():
                match = pattern.match(line.strip())
                if match:
                    rows.append({"date": match[1], "description": match[2], "amount": match[3]})
        return rows


PARSERS = {".csv": CSVParser(), ".xlsx": XLSXParser(), ".pdf": PDFParser()}


def parse_statement(filename, content):
    suffix = Path(filename).suffix.lower()
    parser = PARSERS.get(suffix)
    if not parser:
        raise Clarification("Send a CSV, XLSX or text-extractable PDF statement.")
    try:
        raw = parser.parse(content)
    except (ValueError, csv.Error, KeyError, OSError, RuntimeError) as exc:
        raise Clarification("I couldn't parse that statement. Check the file format and headers.") from exc
    if not raw or len(raw) > 2000:
        raise Clarification("No extractable rows found, or the statement exceeds 2,000 rows.")
    return [normalize_row(row, index) for index, row in enumerate(raw)]


async def duplicate_match(db, user_id, row):
    date = datetime.fromisoformat(row["date"]).replace(tzinfo=UTC)
    candidates = (
        await db.scalars(
            select(Transaction)
            .where(
                Transaction.user_id == user_id,
                Transaction.currency == row["currency"],
                Transaction.amount == Decimal(row["amount"]),
                Transaction.transaction_type == row["transaction_type"],
                Transaction.date >= date - timedelta(days=3),
                Transaction.date < date + timedelta(days=4),
            )
            .limit(50)
        )
    ).all()
    target = row["description"].casefold()
    return next(
        (
            item
            for item in candidates
            if SequenceMatcher(None, target, (item.merchant or item.description).casefold()).ratio() >= 0.72
        ),
        None,
    )


async def preview_statement(db, user_id, filename, content, account_name=None):
    digest = hashlib.sha256(content).hexdigest()
    existing = await db.scalar(
        select(StatementImport).where(
            StatementImport.user_id == user_id,
            StatementImport.file_hash == digest,
            StatementImport.status == "committed",
        )
    )
    if existing:
        raise Clarification("This exact statement was already imported.")
    rows = parse_statement(filename, content)
    account = None
    if account_name:
        account = await db.scalar(
            select(Account).where(Account.user_id == user_id, Account.name == account_name)
        )
        if not account:
            raise Clarification("I couldn't find that account. Create it first or omit the account name.")
    for row in rows:
        if row["status"] == "ready":
            if account and row["currency"] != account.currency:
                row["status"], row["reason"] = "review", "Currency differs from the selected account."
            elif match := await duplicate_match(db, user_id, row):
                row["status"], row["reason"], row["match_id"] = (
                    "duplicate",
                    "Matches a recorded transaction.",
                    match.id,
                )
    item = StatementImport(
        user_id=user_id,
        file_hash=digest,
        filename=filename[:250],
        account_id=account.id if account else None,
        rows={"items": rows},
        status="preview",
        imported_count=0,
    )
    db.add(item)
    await db.flush()
    counts = {
        state: sum(row["status"] == state for row in rows) for state in ("ready", "duplicate", "review")
    }
    sample = "\n".join(
        f"{r['date']} {r['description']}: {r['amount']} {r['currency']}"
        for r in rows
        if r["status"] == "ready"
    )[:800]
    review = "\n".join(
        f"row {row['index'] + 2}: {row.get('reason', 'Needs review')}"
        for row in rows
        if row["status"] == "review"
    )[:400]
    message = (
        f"Statement preview: {filename}\n{len(rows)} rows · {counts['ready']} ready · {counts['duplicate']} duplicates · {counts['review']} need review.\n"
        + (f"Account: {account.name}\n" if account else "Account: unlinked\n")
        + sample
        + ("\nNeeds review (excluded):\n" + review if review else "")
        + "\nConfirm to import ready rows only. Review rows are excluded."
    )
    return item.id, message


async def commit_statement(db, user_id, import_id):
    item = await db.scalar(
        select(StatementImport)
        .where(StatementImport.user_id == user_id, StatementImport.id == import_id)
        .with_for_update()
    )
    if not item or item.status != "preview":
        raise Clarification("That statement preview is no longer available.")
    imported = 0
    for row in item.rows["items"]:
        if row["status"] != "ready" or await duplicate_match(db, user_id, row):
            continue
        source = f"statement:{item.id}:{row['index']}"
        exists = await db.scalar(
            select(Transaction.id).where(
                Transaction.user_id == user_id, Transaction.source_message_id == source
            )
        )
        if exists:
            continue
        transaction = Transaction(
            user_id=user_id,
            date=datetime.fromisoformat(row["date"]).replace(tzinfo=UTC),
            amount=Decimal(row["amount"]),
            currency=row["currency"],
            merchant=row["description"],
            description=row["description"],
            category="Uncategorized",
            subcategory="",
            transaction_type=row["transaction_type"],
            account_id=item.account_id,
            payment_method="",
            source="statement",
            source_message_id=source,
            confidence=1.0,
            notes="Imported from statement preview",
            original_input=f"Statement {item.filename} row {row['index']}",
        )
        db.add(transaction)
        await db.flush()
        await enqueue_sync(db, user_id, "transaction", transaction.id)
        imported += 1
    item.status, item.imported_count = "committed", imported
    return f"Imported {imported} transactions. Duplicates and uncertain rows were skipped."
