import io
from datetime import timedelta
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.database import AgentOpinion, Decision, Delivery, Reminder, Transaction
from app.reminders import SchedulerWorker
from app.schemas import DecisionSynthesis, SpecialistView


async def test_A_reminder_and_delivery(env, future):
    settings, sessions, model, calendar, service = env
    model.route("reminder.create", title="water the plants", start=future.isoformat())
    reply = await service.handle(123, "a", "Remind me tomorrow at 8 PM to water the plants.")
    assert "Reminder saved" in reply.text and reply.reminder_id
    sent = []

    class Messenger:
        async def send(self, user_id, text, **kwargs):
            sent.append((user_id, text))

    worker = SchedulerWorker(sessions, settings, Messenger(), service.briefing)
    await worker.enqueue(future + timedelta(seconds=1))
    await worker.enqueue(future + timedelta(seconds=2))
    await worker.send_pending(future + timedelta(seconds=3))
    await worker.send_pending(future + timedelta(seconds=4))
    assert len(sent) == 1 and sent[0] == (123, "Reminder: water the plants")
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(Delivery)) == 1
        assert (await db.get(Reminder, reply.reminder_id)).status == "awaiting_completion"


async def test_B_calendar_create(env, future):
    _, _, model, calendar, service = env
    model.route(
        "calendar.create",
        title="Gym",
        start=future.isoformat(),
        end=(future + timedelta(hours=1)).isoformat(),
    )
    reply = await service.handle(123, "b", "Gym Friday at 5 PM.")
    assert "Created: Gym" in reply.text
    assert len(calendar.events) == 1


async def test_C_expense_and_idempotency(env):
    _, sessions, model, _, service = env
    model.route(
        "finance.log_expense",
        amount="48k",
        merchant="Coffee shop",
        account="BCA",
        category="Food & Drinks",
        subcategory="Coffee",
    )
    reply = await service.handle(123, "c", "Spent 48k on coffee using BCA.")
    assert "Logged Rp48,000" in reply.text and not reply.confirmation_id
    await service.handle(123, "c", "Spent 48k on coffee using BCA.")
    async with sessions() as db:
        values = list((await db.scalars(select(Transaction))).all())
        assert len(values) == 1
        assert values[0].amount == Decimal(48000)
        assert values[0].original_input == "Spent 48k on coffee using BCA."
        assert values[0].subcategory == "Coffee"


async def test_D_receipt_duplicate_override(env, receipt_image):
    _, sessions, _, _, service = env
    first = await service.handle(123, "d1", "", receipt_image, "file-unique")
    assert "Logged Rp78,000" in first.text
    duplicate = await service.handle(123, "d2", "", receipt_image, "file-unique")
    assert "appears to match" in duplicate.text and duplicate.confirmation_id
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(Transaction)) == 1
    approved = await service.confirm(123, duplicate.confirmation_id, True)
    assert "Logged" in approved.text
    await service.confirm(123, duplicate.confirmation_id, True)
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(Transaction)) == 2


async def test_E_filtered_finance_query(env):
    _, _, model, _, service = env
    model.route("finance.log_expense", amount="48k", category="Food & Drinks")
    await service.handle(123, "e1", "48k coffee")
    model.route("finance.log_expense", amount="100k", category="Transport")
    await service.handle(123, "e2", "100k petrol")
    model.route("finance.query", category="Food & Drinks")
    reply = await service.handle(123, "e3", "How much did I spend on food this month?")
    assert "expenses Rp48,000" in reply.text and "Rp148,000" not in reply.text


async def test_F_excel_sent_as_document(env):
    _, _, _, _, service = env
    reply = await service.handle(123, "f", "/export_finance")
    assert reply.filename.endswith(".xlsx")
    workbook = load_workbook(io.BytesIO(reply.document))
    assert workbook.sheetnames == [
        "Dashboard",
        "Transactions",
        "Categories",
        "Monthly Summary",
        "Accounts",
        "Investments",
        "Net Worth",
        "Decision History",
    ]
    assert workbook["Dashboard"]["B4"].value.startswith("=COUNTA")


async def test_G_council_rounds_persistence(env):
    _, sessions, model, _, service = env
    model.route("decision.request", decision_type="purchase")
    reply = await service.handle(123, "g", "Should I spend Rp12m on a new PC?")
    assert "Confirm your reserves before buying." in reply.text
    assert "Affordability needs to be clear" in reply.text
    assert "Utility may justify" in reply.text
    assert "Confidence:" not in reply.text and "Council:" not in reply.text
    calls = [c for c in model.calls if c[0] is SpecialistView]
    assert len(calls) == 2
    assert len([c for c in model.calls if c[0] is DecisionSynthesis]) == 1
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(AgentOpinion)) == 2
        decision = await db.scalar(select(Decision))
        assert "this_month" in decision.context_snapshot
        assert (
            decision.final_recommendation["common_ground"]
            == "Affordability needs to be clear before you decide."
        )


async def test_H_decision_retrieval(env):
    _, _, model, _, service = env
    model.route("decision.request", decision_type="purchase")
    await service.handle(123, "h1", "Should I buy a PC?")
    model.route("decision.history", search="PC")
    reply = await service.handle(123, "h2", "What did the council previously say about buying a PC?")
    assert "Confirm your reserves" in reply.text
    other = await service.handle(456, "h3", "What about PC?")
    assert "No matching" in other.text


async def test_I_calendar_query_timezone(env, future):
    _, _, model, calendar, service = env
    await calendar.create_event(123, "Gym", future, future + timedelta(hours=1), "setup")
    model.route(
        "calendar.query",
        start=future.replace(hour=0).isoformat(),
        end=(future.replace(hour=0) + timedelta(days=1)).isoformat(),
    )
    reply = await service.handle(123, "i", "What's my schedule tomorrow?")
    assert "20:00" in reply.text and "WIB" in reply.text
